# -*- coding: utf-8 -*-
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="B0C4DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_table(table, col_widths, headers, data, header_bg="1F497D", alt_bg="F4F7FA"):
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], 120, 120, 120, 120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        bg_color = alt_bg if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = str(text)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], 80, 80, 100, 100)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx in [0, 2, 4]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(30, 30, 30)

    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

def add_h(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level==1 else 8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12.5 if level==1 else 11.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(31, 73, 125) if level==1 else RGBColor(50, 50, 50)
    return p

def add_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(10.5)
        rb.font.bold = True
    rt = p.add_run(text)
    rt.font.name = "Times New Roman"
    rt.font.size = Pt(10.5)
    return p

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(10.5)
        rb.font.bold = True
    rt = p.add_run(text)
    rt.font.name = "Times New Roman"
    rt.font.size = Pt(10.5)
    return p

def create_report(week_title, date_range, output_filename):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.79)

    # Header Trường / Khoa
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("TRƯỜNG ĐẠI HỌC LẠC HỒNG\nKHOA CÔNG NGHỆ THÔNG TIN - HỌC PHẦN PHÁT TRIỂN ỨNG DỤNG")
    r_inst.font.name = "Times New Roman"
    r_inst.font.size = Pt(10.5)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(80, 80, 80)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(8)
    r_div = p_div.add_run("— OoO —")
    r_div.font.name = "Times New Roman"
    r_div.font.size = Pt(10)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(3)
    r_title = p_title.add_run(f"BÁO CÁO TIẾN ĐỘ THỰC HIỆN ĐỀ TÀI - {week_title.upper()}")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(14.5)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(31, 73, 125)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(10)
    r_sub = p_sub.add_run("Đề tài: NỀN TẢNG QUẢN TRỊ CHUỖI CUNG ỨNG CHỐNG LÃNG PHÍ\n(SUPPLY CHAIN SPOILAGE PREDICTOR)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(192, 0, 0)

    # Info Box
    info_tbl = doc.add_table(rows=1, cols=2)
    info_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(info_tbl, color="B0C4DE", sz="6")
    c1, c2 = info_tbl.rows[0].cells
    set_cell_background(c1, "F4F7FA")
    set_cell_background(c2, "F4F7FA")
    set_cell_margins(c1, 100, 100, 140, 140)
    set_cell_margins(c2, 100, 100, 140, 140)

    p1 = c1.paragraphs[0]
    p1.add_run("Giảng viên hướng dẫn:\n").font.bold = True
    p1.add_run(f"ThS. Lê Minh Nhật\nThời gian: {date_range}")
    for r in p1.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)

    p2 = c2.paragraphs[0]
    p2.add_run("Nhóm sinh viên thực hiện:\n").font.bold = True
    p2.add_run("1. Đỗ Tấn Du - MSSV: 123001364\n2. Đoàn Minh Quân - MSSV: 123000946")
    for r in p2.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(10)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 1. Mục tiêu tuần
    add_h(doc, "1. Mục tiêu đặt ra trong tuần (Khớp mục 4 & 5 Kế hoạch đăng ký)")
    add_bullet(doc, "Nghiệp vụ cốt lõi: ", "Tinh gọn mô hình hệ thống theo góp ý của GVHD Lê Minh Nhật (tập trung mô hình Bán lẻ Retail: Kho tổng DC -> Cửa hàng -> Khách hàng; bán hàng ra = xuất kho FEFO; loại bỏ điều chuyển nội bộ và thủ tục trả hàng phức tạp).")
    add_bullet(doc, "Chuẩn hóa công thức & thiết kế: ", "Xây dựng 3 ngưỡng an toàn khoa học (Ngưỡng cận date theo % vòng đời RSL, Ngưỡng tồn kho tối thiểu ROP, Ngưỡng nguy cơ hư hỏng Spoilage Risk). Hoàn thiện bộ sơ đồ Use Case và DFD Mức 0, Mức 1.")
    add_bullet(doc, "Hiện thực hóa Cơ sở dữ liệu: ", "Thiết kế và nạp CSDL quan hệ chuẩn 3NF vào Microsoft SQL Server (SSMS), viết Stored Procedures, Views cảnh báo và Triggers tự động.")
    add_bullet(doc, "Hạ tầng mã nguồn: ", "Khởi tạo Git, đẩy toàn bộ mã nguồn lên GitHub và kết nối cộng tác giữa các thành viên nhóm.")

    # 2. Công việc hoàn thành
    add_h(doc, "2. Các công việc cụ thể đã hoàn thành trong tuần")
    tbl_tasks = doc.add_table(rows=1, cols=5)
    widths = [0.45, 2.55, 1.2, 1.9, 0.65]
    headers = ["STT", "Nội Dung Công Việc Đã Thực Hiện", "Người Phụ Trách", "Sản Phẩm Đạt Được", "Tiến Độ"]
    data = [
        [
            "1",
            "Phân tích & tinh gọn quy trình bán lẻ theo dặn dò của GVHD. Chuẩn hóa 8 yêu cầu chức năng (FR) và 4 phi chức năng (NFR). Thiết lập bảng hệ số tiêu thụ K (thời tiết, ngày lễ) và 3 ngưỡng an toàn (RSL, ROP, Spoilage DOS > DUE).",
            "Đỗ Tấn Du\nĐoàn Minh Quân",
            "Tài liệu thuyết minh hoàn chỉnh: Docs/BAO_CAO_NGHIEP_VU_VA_KIEN_TRUC_NOP_THAY.docx",
            "100%"
        ],
        [
            "2",
            "Thiết kế Sơ đồ Use Case (9 chức năng phân quyền Staff/Manager) và Sơ đồ luồng dữ liệu DFD Mức 0 (Context) & DFD Mức 1 (5 tiến trình nghiệp vụ, 3 kho dữ liệu cốt lõi).",
            "Đỗ Tấn Du",
            "Bộ ảnh sơ đồ hoàn chỉnh tại Docs/Diagram_Images/ (use_case.png, dfd_0.png, dfd_1.png)",
            "100%"
        ],
        [
            "3",
            "Xây dựng & tối ưu CSDL Microsoft SQL Server (T-SQL): 12 bảng chuẩn 3NF, 4 Views tính hạn dùng & gợi ý nhập hàng động, Stored Procedure xuất kho FEFO tự động, Function tính rủi ro fn_calculate_spoilage_risk, Trigger trg_after_spoilage_insert.",
            "Đoàn Minh Quân\nĐỗ Tấn Du",
            "File kịch bản Database/retail_spoilage_database.sql nạp thành công vào SSMS LocalDB",
            "100%"
        ],
        [
            "4",
            "Viết kịch bản kiểm thử CSDL tự động bằng Python, kết nối và kiểm tra tính toàn vẹn khóa ngoại, logic cảnh báo cận date và cơ chế trừ kho FEFO.",
            "Đỗ Tấn Du",
            "File Database/test_database_runner.py chạy pass 100% các ca kiểm thử",
            "100%"
        ],
        [
            "5",
            "Tổ chức cấu trúc dự án chuẩn modular (Docs, Database, Sample_Data, Backend, Frontend). Khởi tạo Git repository, cấu hình .gitignore và đẩy toàn bộ lên GitHub.",
            "Đỗ Tấn Du",
            "Repo GitHub hoạt động: https://github.com/DoTanDu/supply-chain-spoilage-predictor",
            "100%"
        ]
    ]
    format_table(tbl_tasks, widths, headers, data)

    # 3. Khó khăn & Giải pháp
    add_h(doc, "3. Khó khăn kỹ thuật gặp phải và giải pháp xử lý")
    add_bullet(doc, "Lệch bảng mã Tiếng Việt trên SQL Server: ", "Khi nạp ban đầu bị lỗi hiển thị font dấu tiếng Việt trong SSMS. Nhóm đã xử lý triệt để bằng cách chuyển mã nguồn T-SQL sang UTF-8 with BOM (utf-8-sig) và khai báo tiền tố N'...' cho toàn bộ chuỗi ký tự Unicode.")
    add_bullet(doc, "Cấu hình dịch vụ máy chủ SQL Server: ", "Dịch vụ MSSQL$SQLEXPRESS bị lỗi đường dẫn tệp hệ thống trên máy local. Nhóm đã chuyển hướng cấu hình kết nối trực tiếp vào Microsoft SQL Server LocalDB ((localdb)\\mssqllocaldb) sẵn có, đảm bảo cả 2 thành viên đều thao tác thuận tiện trên SSMS.")
    add_bullet(doc, "Đồng bộ môi trường làm việc nhóm: ", "Để đảm bảo Quân tải về máy là chạy được ngay, nhóm đã đóng gói tài liệu HUONG_DAN_DATABASE.md chi tiết từng bước kết nối SSMS và chia sẻ quyền cộng tác trên GitHub.")

    # 4. Kế hoạch tuần tới
    add_h(doc, "4. Kế hoạch tuần tiếp theo (21/09/2026 – 27/09/2026)")
    add_p(doc, "Theo đúng tiến độ Kế hoạch thực hiện (bắt đầu giai đoạn STT 6 & 7 từ 21/09 đến 18/10), nhóm phân công:")
    add_bullet(doc, "Đỗ Tấn Du (Frontend): ", "Khởi tạo dự án React (Vite) + TailwindCSS trong thư mục Frontend/; xây dựng Layout Dashboard tổng thể (Sidebar, Header, Alert Badge); dựng màn hình Quản lý Lô hàng & Bảng cảnh báo hạn sử dụng (FEFO).")
    add_bullet(doc, "Đoàn Minh Quân (Backend): ", "Khởi tạo dự án Backend API trong thư mục Backend/; cấu hình kết nối CSDL SQL Server LocalDB; viết các RESTful API xác thực người dùng (JWT) và API danh mục sản phẩm, nhập lô hàng.")

    # 5. Đánh giá
    add_h(doc, "5. Tự đánh giá mức độ hoàn thành")
    add_p(doc, "Tiến độ đạt 100% mục tiêu giai đoạn Thiết kế & Khảo sát (STT 4 & 5 trong bảng Kế hoạch thực hiện đề tài). Toàn bộ nền tảng nghiệp vụ, tài liệu học thuật, sơ đồ hệ thống và Cơ sở dữ liệu đã sẵn sàng để bước vào giai đoạn Lập trình Backend - Frontend.")

    # Chữ ký
    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.paragraph_format.space_before = Pt(14)
    r_sig = p_sig.add_run("Đồng Nai, Ngày 19 tháng 09 năm 2026\nĐại diện nhóm sinh viên thực hiện\n\n\n\nĐỗ Tấn Du")
    r_sig.font.name = "Times New Roman"
    r_sig.font.size = Pt(10.5)

    doc.save(output_filename)
    print(f"Generated: {output_filename}")

def main():
    # Tạo báo cáo cho Tuần 4 (14/09/2026 - 20/09/2026)
    create_report("Tuần 4", "14/09/2026 – 20/09/2026", r"c:\thuyet minh\Docs\BAO_CAO_TIEN_DO_TUAN_4.docx")
    # Tạo dự phòng file Tuần 3 (dành cho trường hợp thầy tính theo mốc 3 tuần)
    create_report("Tuần 3 & 4", "07/09/2026 – 20/09/2026", r"c:\thuyet minh\Docs\BAO_CAO_TIEN_DO_TUAN_3.docx")

if __name__ == "__main__":
    main()
