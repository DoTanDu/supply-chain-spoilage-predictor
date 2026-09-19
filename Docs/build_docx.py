# -*- coding: utf-8 -*-
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def format_table(table, col_widths, headers, data, header_bg="1F497D", alt_bg="F2F5F8"):
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    # Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        bg_color = alt_bg if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].text = str(text)
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[col_idx].paragraphs[0]
            if col_idx == 0 or len(str(text)) < 8:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(40, 40, 40)

    # Set column widths
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = Inches(w)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(31, 73, 125) # Navy
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor(54, 96, 145)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(40, 40, 40)
    return p

def add_body_p(doc, text, bold_prefix="", italic_note=""):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Times New Roman"
        r_b.font.size = Pt(11)
        r_b.font.bold = True
        r_b.font.color.rgb = RGBColor(30, 30, 30)
    if text:
        r_t = p.add_run(text)
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(11)
        r_t.font.color.rgb = RGBColor(40, 40, 40)
    if italic_note:
        r_i = p.add_run(italic_note)
        r_i.font.name = "Times New Roman"
        r_i.font.size = Pt(10.5)
        r_i.font.italic = True
        r_i.font.color.rgb = RGBColor(100, 100, 100)
    return p

def add_bullet_p(doc, bold_title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_title:
        rb = p.add_run(bold_title)
        rb.font.name = "Times New Roman"
        rb.font.size = Pt(11)
        rb.font.bold = True
    rt = p.add_run(text)
    rt.font.name = "Times New Roman"
    rt.font.size = Pt(11)
    return p

def add_image_with_caption(doc, img_path, caption_text, width_in=5.8):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(width_in))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        run_cap = p_cap.add_run(caption_text)
        run_cap.font.name = "Times New Roman"
        run_cap.font.size = Pt(10)
        run_cap.font.italic = True
        run_cap.font.bold = True
        run_cap.font.color.rgb = RGBColor(60, 60, 60)

def main():
    doc = Document()

    # Set page margins: Top 2cm, Bottom 2cm, Left 2.5cm, Right 2cm
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.79)

    # ---------------- COVER / HEADER ----------------
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_after = Pt(2)
    r_inst = p_inst.add_run("TRƯỜNG ĐẠI HỌC LẠC HỒNG\nKHOA CÔNG NGHỆ THÔNG TIN - BỘ MÔN KỸ THUẬT PHẦN MỀM")
    r_inst.font.name = "Times New Roman"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(80, 80, 80)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(16)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("BÁO CÁO PHÂN TÍCH NGHIỆP VỤ & THIẾT KẾ KIẾN TRÚC HỆ THỐNG")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(31, 73, 125)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run("Đề tài: NỀN TẢNG QUẢN TRỊ CHUỖI CUNG ỨNG CHỐNG LÃNG PHÍ\n(SUPPLY CHAIN SPOILAGE PREDICTOR)")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(192, 0, 0)

    # Student & Teacher table / info box
    info_table = doc.add_table(rows=0, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(info_table, color="B0C4DE", sz="6")
    row_cells = info_table.add_row().cells
    set_cell_background(row_cells[0], "F4F7FA")
    set_cell_background(row_cells[1], "F4F7FA")
    set_cell_margins(row_cells[0], 120, 120, 160, 160)
    set_cell_margins(row_cells[1], 120, 120, 160, 160)

    p1 = row_cells[0].paragraphs[0]
    p1.add_run("Giảng viên hướng dẫn:\n").font.bold = True
    p1.add_run("ThS. Lê Minh Nhật\nEmail: nhatlm@lhu.edu.vn\nSĐT: 0389571228")
    for r in p1.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(10.5)

    p2 = row_cells[1].paragraphs[0]
    p2.add_run("Nhóm sinh viên thực hiện:\n").font.bold = True
    p2.add_run("1. Đỗ Tấn Du - MSSV: 123001364\n2. Đoàn Minh Quân - MSSV: 123000946\nHọc phần: Phát triển ứng dụng")
    for r in p2.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(10.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ---------------- PHẦN 1 ----------------
    add_heading_1(doc, "I. PHÂN TÍCH NGHIỆP VỤ HỆ THỐNG (BUSINESS ANALYSIS)")
    
    add_heading_2(doc, "1.1. Bối cảnh thực tế và Phạm vi bài toán")
    add_body_p(doc, "Trong lĩnh vực bán lẻ thực phẩm và hàng tiêu dùng nhanh (FMCG) như các chuỗi cửa hàng tiện lợi, bách hóa (Bách Hóa Xanh, WinMart, GS25, Circle K), việc kiểm soát hàng tồn kho và hạn sử dụng (Shelf-life) là vấn đề sống còn. Nếu ước lượng nhập hàng theo cảm tính thủ công, cửa hàng sẽ đối mặt với 2 rủi ro nghiêm trọng: (1) Nhập dư dẫn đến hàng quá date, hư hỏng, phải tiêu hủy gây tổn thất chi phí trực tiếp; (2) Nhập thiếu làm đứt gãy nguồn cung, mất doanh thu và giảm sự hài lòng của khách hàng.")
    add_body_p(doc, "Dựa trên định hướng chuẩn hóa của Giảng viên hướng dẫn, hệ thống xác định rõ phạm vi nghiệp vụ tập trung vào bài toán Kho bán lẻ (Retail Store): Tiếp nhận hàng hóa theo đợt từ Kho trung tâm (DC), quản lý hạn dùng từng lô, bán lẻ trực tiếp ra cho người tiêu dùng theo nguyên tắc FEFO, loại bỏ hoàn toàn các quy trình luân chuyển ngang giữa các shop và thủ tục trả hàng lỗi phức tạp để tập trung tối đa vào bài toán Chống lãng phí (Spoilage Prevention).")

    add_heading_2(doc, "1.2. Các quy trình nghiệp vụ cốt lõi")
    add_bullet_p(doc, "1. Quy trình Tiếp nhận Lô hàng từ Kho tổng (Goods Receipt): ", "Hàng hóa từ Kho tổng chuyển về shop được nhân viên tiếp nhận và nhập vào hệ thống theo Lô hàng (Batch). Mỗi lô bắt buộc lưu thông tin: Mã lô (Batch Code), Ngày sản xuất (MFG), Hạn sử dụng (EXP), Số lượng nhập và Giá nhập.")
    add_bullet_p(doc, "2. Quy trình Bán hàng lẻ & Trừ tồn kho tự động (FEFO Sales Out): ", "Khi phát sinh giao dịch bán hàng (hoặc import dữ liệu bán từ file Excel/CSV), hệ thống tự động tìm và trừ tồn kho ở lô hàng có hạn sử dụng gần nhất (First Expired, First Out). Cơ chế này đảm bảo hàng cũ luôn được ưu tiên bán trước, triệt tiêu tình trạng hàng mới về đè lên hàng cũ dẫn đến hàng cũ bị quá date.")
    add_bullet_p(doc, "3. Quy trình Kiểm soát hạn sử dụng & Cảnh báo Spoilage: ", "Hệ thống tự động quét hạn sử dụng hàng ngày và phân loại 3 cấp độ cảnh báo trực quan: Cảnh báo Vàng (còn <= 7 ngày - chú ý theo dõi), Cảnh báo Đỏ (còn <= 3 ngày - nguy cơ cao, cần dán tem giảm giá xả hàng gấp), Khóa hết hạn (còn <= 0 ngày - khóa mã vạch, cấm bán).")
    add_bullet_p(doc, "4. Quy trình Dự báo Tiêu thụ & Đề xuất Đặt hàng (Reorder Forecasting): ", "Dựa vào dữ liệu lịch sử bán hàng kết hợp các biến ngoại cảnh (thời tiết nắng/mưa, ngày cuối tuần, ngày lễ), hệ thống tính toán Tốc độ bán trung bình ngày (d) và Điểm đặt hàng lại (Reorder Point - ROP) để tự động xuất danh sách đề xuất số lượng cần xin Kho tổng cấp hàng.")
    add_bullet_p(doc, "5. Quy trình Tiêu hủy hàng hỏng/quá date (Spoilage Disposal Log): ", "Đối với các sản phẩm quá date hoặc hư hỏng thực tế trong quá trình trưng bày, nhân viên tạo phiếu tiêu hủy. Hệ thống lập tức trừ tồn kho về 0, chuyển trạng thái lô thành ĐÃ HỦY và lưu vào nhật ký lãng phí kèm giá trị thiệt hại tài chính để vẽ biểu đồ thống kê.")

    add_heading_2(doc, "1.3. Phân tích các yếu tố ảnh hưởng đến tiêu thụ & Bảng hệ số tác động (K)")
    add_body_p(doc, "Hệ thống không dự báo đơn thuần bằng giá trị bình quân phẳng mà tích hợp mô hình trọng số ngoại cảnh thông qua Hệ số tác động tiêu thụ K (Impact Multiplier):")

    tbl_k = doc.add_table(rows=1, cols=4)
    k_widths = [1.2, 1.8, 1.2, 2.3]
    k_headers = ["Yếu tố ngoại cảnh", "Nhóm sản phẩm chịu ảnh hưởng", "Hệ số K", "Tác động nghiệp vụ"]
    k_data = [
        ["Ngày thường", "Tất cả các ngành hàng", "K = 1.00", "Sức mua ổn định theo định mức cơ bản."],
        ["Cuối tuần (T7, CN)", "Bia, nước ngọt, đồ ăn vặt, đồ tươi", "K = 1.25", "Nhu cầu giải trí, liên hoan tăng 25%."],
        ["Dịp Lễ / Tết", "Bánh kẹo, quà tặng, thực phẩm chế biến", "K = 1.60", "Sức mua tăng vọt 60%, cần tăng ROP."],
        ["Trời nắng nóng (>34°C)", "Nước giải khát, nước khoáng, kem, sữa chua", "K = 1.40", "Tiêu thụ đồ mát tăng 40%, hạn dùng ngắn."],
        ["Trời mưa to / Bão", "Toàn bộ cửa hàng (trừ đồ hộp, mì ăn liền)", "K = 0.75", "Lượng khách ghé giảm 25%, giảm nhập đồ tươi."]
    ]
    format_table(tbl_k, k_widths, k_headers, k_data)
    add_body_p(doc, "", italic_note="* Công thức dự báo điều chỉnh: Nhu cầu dự báo = Nhu cầu bán trung bình ngày x K x Thời gian cung ứng (Lead time).")

    add_heading_2(doc, "1.4. Bảng Yêu cầu Hệ thống (System Requirements)")
    add_heading_3(doc, "A. Yêu cầu chức năng (Functional Requirements - FR)")
    
    tbl_fr = doc.add_table(rows=1, cols=3)
    fr_widths = [1.0, 1.8, 3.7]
    fr_headers = ["Mã Yêu Cầu", "Tên Chức Năng", "Mô Tả Chi Tiết Nghiệp Vụ"]
    fr_data = [
        ["FR-01", "Quản lý Danh mục & Sản phẩm", "CRUD thông tin sản phẩm, mã vạch SKU, đơn vị tính, hạn sử dụng chuẩn, ngưỡng tồn min/max."],
        ["FR-02", "Tiếp nhận Lô hàng từ Kho tổng", "Nhập thông tin lô hàng (Mã lô, NSX, HSD, số lượng nhập, giá nhập), kích hoạt trạng thái ACTIVE."],
        ["FR-03", "Bán hàng & Trừ kho FEFO", "Ghi nhận hóa đơn bán lẻ, hệ thống tự động quét và trừ số lượng của lô có HSD gần nhất."],
        ["FR-04", "Import Dữ liệu Lịch sử (Excel/CSV)", "Tải file Excel lịch sử giao dịch bán hàng (ngày, SKU, số lượng, thời tiết, ngày lễ) để nạp dữ liệu nhanh."],
        ["FR-05", "Cảnh báo Spoilage Đa cấp", "Phát cảnh báo Vàng (còn <= 7 ngày), Đỏ (còn <= 3 ngày), và tự động khóa mã khi sản phẩm hết hạn."],
        ["FR-06", "Dự báo & Đề xuất Đặt hàng", "Tính tốc độ tiêu thụ d và ROP, kết hợp hệ số thời tiết/lễ để tự động xuất danh sách xin cấp hàng."],
        ["FR-07", "Ghi nhận Tiêu hủy Hàng hỏng", "Lập phiếu hủy lô quá date, cập nhật tồn kho về 0, lưu lý do và tính toán giá trị thiệt hại tài chính."],
        ["FR-08", "Dashboard Báo cáo & Thống kê", "Trực quan hóa biểu đồ doanh thu, top sản phẩm nguy cơ hết hạn cao, và tỷ lệ hao hụt hàng hóa."]
    ]
    format_table(tbl_fr, fr_widths, fr_headers, fr_data)

    add_heading_3(doc, "B. Yêu cầu phi chức năng (Non-Functional Requirements - NFR)")
    tbl_nfr = doc.add_table(rows=1, cols=3)
    nfr_widths = [1.0, 1.8, 3.7]
    nfr_headers = ["Mã NFR", "Tiêu Chí", "Chỉ Số Đạt Chuẩn"]
    nfr_data = [
        ["NFR-01", "Hiệu năng (Performance)", "Tốc độ tải trang < 1.5s; Import file Excel 1.000 dòng hoàn tất dưới 3s."],
        ["NFR-02", "Độ chính xác tồn kho", "Cập nhật tồn kho Real-time sau mỗi giao dịch; tuyệt đối không để xảy ra tồn kho âm."],
        ["NFR-03", "Trải nghiệm UI/UX", "Giao diện Responsive chạy tốt trên Desktop và Tablet; màu sắc cảnh báo chuẩn (Xanh - Vàng - Đỏ)."],
        ["NFR-04", "Bảo mật & Phân quyền", "Xác thực JWT Token, phân chia chặt chẽ vai trò Nhân viên (Staff) và Cửa hàng trưởng (Manager)."]
    ]
    format_table(tbl_nfr, nfr_widths, nfr_headers, nfr_data)

    # ---------------- PHẦN 2 ----------------
    add_heading_1(doc, "II. THIẾT KẾ SƠ ĐỒ USE CASE & SƠ ĐỒ LUỒNG DỮ LIỆU (DFD)")

    add_heading_2(doc, "2.1. Phân quyền Tác nhân (Actors)")
    add_bullet_p(doc, "Nhân viên / Thủ kho (Staff): ", "Thực hiện các tác vụ tác nghiệp tại quầy: Nhập lô hàng từ kho tổng, bán hàng lẻ trừ kho, theo dõi cảnh báo date hàng ngày, và tạo phiếu hủy hàng quá date.")
    add_bullet_p(doc, "Cửa hàng trưởng / Admin (Store Manager): ", "Toàn quyền quản lý danh mục sản phẩm, cấu hình ngưỡng cảnh báo, import file Excel lịch sử giao dịch, xem Dashboard phân tích doanh thu và duyệt đơn đề xuất đặt hàng gửi lên Kho tổng.")

    add_heading_2(doc, "2.2. Sơ đồ Use Case của Hệ thống")
    add_image_with_caption(doc, r"c:\thuyet minh\Docs\Diagram_Images\use_case.png", "Hình 1: Sơ đồ Use Case Hệ thống Quản trị Bán lẻ Chống lãng phí (Chi tiết 6 phân hệ tác nghiệp)")

    add_body_p(doc, "Bảng đặc tả tóm tắt các Use Case chính trong hệ thống:")
    tbl_uc = doc.add_table(rows=1, cols=3)
    uc_widths = [1.8, 1.2, 3.5]
    uc_headers = ["Tên Use Case", "Tác Nhân Chính", "Mô Tả Tóm Tắt Luồng Xử Lý"]
    uc_data = [
        ["Đăng nhập hệ thống (JWT)", "Staff, Manager", "Xác thực tài khoản và trả về JWT Token kèm phân quyền tương ứng."],
        ["Tiếp nhận & Nhập lô từ Kho tổng", "Staff", "Kiểm đếm HSD thực tế, lưu thông tin lô, kích hoạt Shelf-life, in tem nhãn kệ."],
        ["Bán hàng lẻ / Trừ kho FEFO", "Staff", "Quét mã vạch POS, tự động quét chọn lô HSD gần nhất trừ tồn, lưu lịch sử và biến thời tiết/lễ."],
        ["Giám sát HSD & Cảnh báo Spoilage", "Staff, Manager", "Quét tự động đếm lùi ngày HSD, phân loại 3 mức RSL, đảo hàng, dán tem giảm giá, khóa mã khi hết hạn."],
        ["Quy trình Tiêu hủy hàng hỏng", "Staff, Manager", "Nhân viên lập phiếu hủy, Quản lý duyệt, hệ thống trừ tồn về 0 và hạch toán thiệt hại tài chính."],
        ["Dự báo & Đề xuất Đặt hàng", "Manager, DC", "Phân tích tốc độ bán d và hệ số K, tính ROP, cảnh báo DOS>DUE, sinh đề xuất và gửi đơn về Kho tổng."],
        ["Quản lý Danh mục & Cấu hình", "Manager", "CRUD sản phẩm, cấu hình số ngày cảnh báo cận date riêng từng nhóm hàng và tồn kho an toàn."],
        ["Import Excel lịch sử bán", "Manager", "Đọc file Excel bán hàng và nạp hàng loạt vào kho dữ liệu bán hàng."],
        ["Dashboard Thống kê & Báo cáo", "Manager", "Xem biểu đồ doanh thu, tỷ lệ hao hụt hàng hóa (Spoilage Rate %) và xuất báo cáo kiểm kê định kỳ."]
    ]
    format_table(tbl_uc, uc_widths, uc_headers, uc_data)

    add_heading_2(doc, "2.3. Sơ đồ Luồng Dữ liệu DFD Mức 0 (Sơ đồ ngữ cảnh)")
    add_image_with_caption(doc, r"c:\thuyet minh\Docs\Diagram_Images\dfd_0.png", "Hình 2: Sơ đồ Luồng Dữ liệu DFD Mức 0 (Context Diagram)")
    add_body_p(doc, "Thuyết minh luồng dữ liệu mức ngữ cảnh:")
    add_bullet_p(doc, "Tác nhân Người dùng (Nhân viên & Quản lý): ", "Gửi vào hệ thống thông tin lô hàng nhập, giao dịch bán lẻ hoặc file Excel bán hàng; nhận lại từ hệ thống các cảnh báo cận date, danh sách tồn kho và báo cáo tỷ lệ hao hụt.")
    add_bullet_p(doc, "Thực thể Kho tổng (DC): ", "Cung cấp hàng hóa kèm thông tin hạn sử dụng về cho hệ thống cửa hàng; nhận về Đơn đề xuất nhập hàng tối ưu do hệ thống tự động tính toán.")

    add_heading_2(doc, "2.4. Sơ đồ Luồng Dữ liệu DFD Mức 1 (Phân rã chức năng)")
    add_image_with_caption(doc, r"c:\thuyet minh\Docs\Diagram_Images\dfd_1.png", "Hình 3: Sơ đồ Luồng Dữ liệu DFD Mức 1 chi tiết 5 tiến trình và 3 kho dữ liệu")
    add_body_p(doc, "Thuyết minh chi tiết 5 tiến trình xử lý và 3 kho lưu trữ dữ liệu:")
    add_bullet_p(doc, "Tiến trình 1.0 (Nhập lô từ Kho tổng): ", "Nhận thông tin lô và HSD từ Kho tổng và xác nhận từ nhân viên, tiến hành ghi bản ghi lô mới vào Kho D1 (Sản phẩm & Lô hàng).")
    add_bullet_p(doc, "Tiến trình 2.0 (Bán hàng lẻ theo FEFO): ", "Khi có giao dịch bán hoặc upload Excel, tiến trình truy vấn Kho D1 để tìm lô có HSD gần nhất để trừ tồn kho; đồng thời ghi lịch sử giao dịch bán vào Kho D2.")
    add_bullet_p(doc, "Tiến trình 3.0 (Quét HSD & Bắn cảnh báo): ", "Đọc dữ liệu hạn sử dụng từ Kho D1, tính toán số ngày còn lại và phát cảnh báo Đỏ/Vàng/Hết hạn đến giao diện người dùng.")
    add_bullet_p(doc, "Tiến trình 4.0 (Dự báo Tiêu thụ & Đề xuất Đặt hàng): ", "Phân tích dữ liệu bán quá khứ từ Kho D2 kết hợp hệ số thời tiết/lễ, đối chiếu với tồn kho thực tế từ Kho D1 để tính điểm ROP, xuất đề xuất cho Quản lý và gửi đơn về Kho tổng DC.")
    add_bullet_p(doc, "Tiến trình 5.0 (Xử lý Tiêu hủy hàng quá hạn): ", "Tiếp nhận lệnh hủy từ nhân viên, cập nhật số lượng tồn của lô về 0 trong Kho D1, đồng thời lưu chi tiết lý do và giá trị thiệt hại vào Kho D3 (Nhật ký Hủy hàng).")

    # ---------------- PHẦN 3 ----------------
    add_heading_1(doc, "III. KIẾN TRÚC TỔNG THỂ CỦA DỰ ÁN (OVERALL ARCHITECTURE)")
    add_body_p(doc, "Hệ thống được thiết kế theo mô hình Kiến trúc 3 tầng phân tán (3-Tier Client-Server Architecture) kết nối thông qua giao thức chuẩn RESTful API:")

    tbl_arch = doc.add_table(rows=1, cols=3)
    arch_widths = [1.5, 1.8, 3.2]
    arch_headers = ["Tầng Kiến Trúc", "Công Nghệ Đại Diện", "Vai Trò & Nhiệm Vụ"]
    arch_data = [
        ["1. Presentation Layer (Tầng Trình Diễn)", "React.js (SPA), Vite, TailwindCSS, Chart.js", "Giao diện Web tương tác trực quan cho người dùng, xử lý render bảng biểu, biểu đồ Dashboard và đọc file Excel ở Client."],
        ["2. Application Layer (Tầng Nghiệp Vụ)", "Node.js (Express) / Python (FastAPI), JWT, Cron", "Cung cấp hệ thống RESTful API, thực thi thuật toán trừ kho FEFO, thuật toán dự báo ROP và tác vụ quét HSD định kỳ."],
        ["3. Data Layer (Tầng Dữ Liệu)", "PostgreSQL / MySQL 8.0, ORM (Prisma/TypeORM/SQLAlchemy)", "Lưu trữ dữ liệu có cấu trúc, đảm bảo tính toàn vẹn quan hệ (Relational Integrity) và giao dịch nguyên tử (ACID)."]
    ]
    format_table(tbl_arch, arch_widths, arch_headers, arch_data)

    add_body_p(doc, "Ưu điểm vượt trội của kiến trúc này:")
    add_bullet_p(doc, "Tách biệt độc lập (Decoupled): ", "Frontend và Backend giao tiếp hoàn toàn qua chuẩn JSON API. Đỗ Tấn Du (Frontend) và Đoàn Minh Quân (Backend) có thể code song song dựa trên bản tài liệu API Contract mà không bị phụ thuộc tiến độ.")
    add_bullet_p(doc, "Dễ bảo trì và mở rộng: ", "Sau này khi cần nâng cấp thuật toán Machine Learning hoặc phát triển ứng dụng di động (Mobile App) cho thủ kho quét mã vạch, tầng Backend hoàn toàn tái sử dụng được 100%.")

    # ---------------- PHẦN 4 ----------------
    add_heading_1(doc, "IV. KIẾN TRÚC BACKEND & FRONTEND")

    add_heading_2(doc, "4.1. Kiến trúc Backend (Backend Architecture)")
    add_body_p(doc, "Backend được tổ chức theo mô hình Kiến trúc phân lớp (Layered Architecture) chuẩn công nghiệp:")
    add_bullet_p(doc, "Controller Layer: ", "Đón nhận các HTTP Request (GET, POST, PUT, DELETE), xác thực quyền truy cập thông qua JWT Middleware và kiểm tra tính hợp lệ của dữ liệu đầu vào (Validation).")
    add_bullet_p(doc, "Service Layer (Trọng tâm): ", "Nơi cài đặt toàn bộ logic nghiệp vụ đặc thù: Xử lý giao dịch trừ kho FEFO (khóa bản ghi tránh Race Condition), chạy hàm tính toán Spoilage Risk Index, tính điểm đặt hàng ROP.")
    add_bullet_p(doc, "Repository / Data Access Layer: ", "Đóng gói các câu lệnh truy vấn SQL thông qua ORM, thao tác với CSDL một cách an toàn, chống lỗ hổng SQL Injection.")
    add_bullet_p(doc, "Scheduled Worker (Cron Job): ", "Tiến trình chạy ngầm tự động vào lúc 00:00 hàng ngày để quét kiểm tra toàn bộ lô hàng, tự động chuyển trạng thái các lô hết hạn sang EXPIRED và gửi danh sách cảnh báo.")

    add_heading_2(doc, "4.2. Kiến trúc Frontend (Frontend Architecture)")
    add_body_p(doc, "Frontend được xây dựng theo mô hình Single Page Application (SPA) kết hợp Kiến trúc hướng thành phần (Component-Based Architecture):")
    add_bullet_p(doc, "Component-Based UI: ", "Giao diện chia nhỏ thành các module độc lập: BatchTable (bảng lô hàng có highlight màu), SpoilageAlertCard (thẻ cảnh báo), MetricKPI (thẻ thống kê), ExcelUploader (nút upload và preview file).")
    add_bullet_p(doc, "State Management (Quản lý trạng thái): ", "Sử dụng React Context API hoặc Zustand để đồng bộ dữ liệu người dùng đăng nhập và danh sách cảnh báo tức thì xuyên suốt các trang.")
    add_bullet_p(doc, "Client-Side Excel Processing: ", "Sử dụng thư viện SheetJS (xlsx) để đọc và kiểm tra lỗi cú pháp file Excel ngay trên trình duyệt trước khi gửi dữ liệu sạch về Backend, giảm tải cho máy chủ.")

    # ---------------- PHẦN 5 ----------------
    add_heading_1(doc, "V. THIẾT KẾ CƠ SỞ DỮ LIỆU (DATABASE SELECTION & DESIGN)")

    add_heading_2(doc, "5.1. Lựa chọn Hệ quản trị CSDL & Luận cứ khoa học")
    add_body_p(doc, "Nhóm đề xuất sử dụng PostgreSQL (hoặc MySQL 8.0) - Hệ quản trị cơ sở dữ liệu quan hệ (RDBMS). Các lý do khoa học bao gồm:")
    add_bullet_p(doc, "Tuân thủ tuyệt đối chuẩn ACID (Atomicity, Consistency, Isolation, Durability): ", "Nghiệp vụ kho bán lẻ không được phép xảy ra sai sót. Khi bán 1 đơn hàng, việc trừ tồn kho lô và ghi lịch sử bán phải thành công đồng thời trong 1 Transaction; nếu xảy ra lỗi, hệ thống tự động Rollback, không bao giờ để xảy ra tình trạng mất hàng hoặc tồn kho âm.")
    add_bullet_p(doc, "Ràng buộc quan hệ 1 - Nhiều chặt chẽ: ", "Cấu trúc dữ liệu có quan hệ liên kết phân cấp rất rõ ràng: 1 Sản phẩm có Nhiều Lô hàng; 1 Lô hàng có Nhiều Giao dịch xuất/hủy. RDBMS xử lý các phép JOIN và khóa ngoại (Foreign Key) tối ưu và toàn vẹn hơn nhiều so với NoSQL.")

    add_heading_2(doc, "5.2. Danh sách và Cấu trúc 5 Bảng CSDL cốt lõi")

    add_body_p(doc, "Bảng 1: users (Quản lý tài khoản & Phân quyền)", bold_prefix="")
    tbl_u = doc.add_table(rows=1, cols=4)
    format_table(tbl_u, [1.5, 1.2, 1.0, 2.8], ["Tên Cột", "Kiểu Dữ Liệu", "Khóa", "Mô Tả"], [
        ["id", "INT", "PK", "Mã định danh người dùng tự tăng."],
        ["username", "VARCHAR(50)", "UNIQUE", "Tên đăng nhập hệ thống."],
        ["password_hash", "VARCHAR(255)", "", "Mật khẩu đã mã hóa an toàn (bcrypt)."],
        ["role", "VARCHAR(20)", "", "Vai trò: 'STAFF' (Nhân viên) hoặc 'MANAGER' (Quản lý)."]
    ])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_body_p(doc, "Bảng 2: products (Danh mục thông tin sản phẩm)", bold_prefix="")
    tbl_p = doc.add_table(rows=1, cols=4)
    format_table(tbl_p, [1.5, 1.2, 1.0, 2.8], ["Tên Cột", "Kiểu Dữ Liệu", "Khóa", "Mô Tả"], [
        ["id", "INT", "PK", "Mã sản phẩm tự tăng."],
        ["sku", "VARCHAR(50)", "UNIQUE", "Mã vạch / Mã SKU của sản phẩm."],
        ["name", "VARCHAR(150)", "", "Tên thương phẩm hàng hóa."],
        ["category", "VARCHAR(50)", "", "Loại ngành hàng (Đồ tươi, Nước giải khát, Đồ khô...)."],
        ["unit", "VARCHAR(20)", "", "Đơn vị tính (Hộp, Chai, Gói, Kg)."],
        ["min_stock", "INT", "", "Ngưỡng tồn kho an toàn tối thiểu (ROP cơ bản)."],
        ["warning_days", "INT", "", "Số ngày cấu hình cảnh báo cận date riêng của sản phẩm."]
    ])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_body_p(doc, "Bảng 3: batches (Lô hàng tại cửa hàng - Cốt lõi bài toán Date)", bold_prefix="")
    tbl_b = doc.add_table(rows=1, cols=4)
    format_table(tbl_b, [1.5, 1.2, 1.0, 2.8], ["Tên Cột", "Kiểu Dữ Liệu", "Khóa", "Mô Tả"], [
        ["id", "INT", "PK", "Mã định danh lô hàng."],
        ["batch_code", "VARCHAR(50)", "UNIQUE", "Mã lô thực tế trên bao bì (VD: LOT-2026-001)."],
        ["product_id", "INT", "FK", "Liên kết với bảng products(id)."],
        ["import_date", "DATE", "", "Ngày nhập hàng từ Kho tổng về shop."],
        ["expiry_date", "DATE", "", "Hạn sử dụng của lô (Cột quan trọng nhất)."],
        ["quantity", "INT", "", "Số lượng tồn kho hiện tại của lô trên kệ."],
        ["status", "VARCHAR(20)", "", "Trạng thái: 'ACTIVE', 'WARNING', 'EXPIRED', 'DISPOSED'."]
    ])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_body_p(doc, "Bảng 4: sales_history (Lịch sử bán hàng kèm biến ngoại cảnh)", bold_prefix="")
    tbl_s = doc.add_table(rows=1, cols=4)
    format_table(tbl_s, [1.5, 1.2, 1.0, 2.8], ["Tên Cột", "Kiểu Dữ Liệu", "Khóa", "Mô Tả"], [
        ["id", "INT", "PK", "Mã giao dịch bán."],
        ["product_id", "INT", "FK", "Sản phẩm được bán."],
        ["sale_date", "DATE", "", "Ngày phát sinh giao dịch bán."],
        ["quantity_sold", "INT", "", "Số lượng sản phẩm bán ra."],
        ["weather", "VARCHAR(30)", "", "Thời tiết ghi nhận: 'SUNNY', 'RAINY', 'NORMAL'."],
        ["is_holiday", "BOOLEAN", "", "Đánh dấu ngày Lễ hoặc Cuối tuần (True/False)."]
    ])

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_body_p(doc, "Bảng 5: spoilage_records (Nhật ký hàng tiêu hủy - Đo lường lãng phí)", bold_prefix="")
    tbl_sp = doc.add_table(rows=1, cols=4)
    format_table(tbl_sp, [1.5, 1.2, 1.0, 2.8], ["Tên Cột", "Kiểu Dữ Liệu", "Khóa", "Mô Tả"], [
        ["id", "INT", "PK", "Mã biên bản tiêu hủy."],
        ["batch_id", "INT", "FK", "Lô hàng bị tiêu hủy."],
        ["disposed_date", "DATE", "", "Ngày thực hiện tiêu hủy."],
        ["quantity", "INT", "", "Số lượng sản phẩm bị vứt bỏ."],
        ["cost_loss", "DECIMAL(12,2)", "", "Số tiền thiệt hại tài chính do hàng hỏng (VNĐ)."],
        ["reason", "VARCHAR(255)", "", "Lý do: 'Quá date', 'Móp méo rách bao bì', 'Hư hỏng'."]
    ])

    # ---------------- PHẦN 6 ----------------
    add_heading_1(doc, "VI. ĐỀ XUẤT TÍCH HỢP AI AGENT SKILLS (THEO TÀI LIỆU GVHD CHIA SẺ)")
    add_body_p(doc, "Căn cứ vào kho thư viện mã nguồn mở Agentic Awesome Skills (https://github.com/sickn33/agentic-awesome-skills) do ThS. Lê Minh Nhật chia sẻ, nhóm nhận thấy đây là cơ hội tuyệt vời để nâng tầm đề tài từ một Web quản lý thông thường thành một Ứng dụng thông minh có trợ lý AI (AI-Powered Retail System).")
    add_body_p(doc, "Nhóm đề xuất nghiên cứu và tích hợp 3 kỹ năng AI Agent phù hợp nhất vào hệ thống:")
    add_bullet_p(doc, "1. Skill 'Data Analysis & Anomaly Detection' (Tác nhân Phân tích Tiêu thụ): ", "Agent tự động đọc lịch sử bán hàng và phát hiện các biến động bất thường (Ví dụ: sản phẩm bán chậm đột ngột trong tuần qua dù trời nắng) để đưa ra cảnh báo sớm nguy cơ ế hàng trước khi chạm date.")
    add_bullet_p(doc, "2. Skill 'Executive Summary & Natural Language Ordering' (Tác nhân Đề xuất Đơn hàng): ", "Thay vì chỉ hiển thị số liệu khô khan, Agent sẽ tự động soạn thảo một bản tóm tắt đề xuất hoàn chỉnh bằng ngôn ngữ tự nhiên: 'Dự báo cuối tuần này có mưa bão, khuyến nghị tăng nhập 30% mì ăn liền và giảm 20% nước giải khát. Cần đặt thêm 40 hộp Sữa tươi LOT-03 từ Kho tổng'.")
    add_bullet_p(doc, "3. Skill 'Daily Retail Assistant' (Trợ lý Hỏi đáp Tồn kho & Date): ", "Tích hợp Chatbot cho phép Cửa hàng trưởng dùng câu hỏi tự nhiên như: 'Hôm nay có lô nào sắp hết hạn không?' hoặc 'Mặt hàng nào đang có nguy cơ lãng phí cao nhất?' và nhận câu trả lời tức thì kèm bảng số liệu trích xuất từ CSDL.")

    # Save document
    output_path = r"c:\thuyet minh\Docs\BAO_CAO_NGHIEP_VU_VA_KIEN_TRUC_NOP_THAY.docx"
    doc.save(output_path)
    print(f"Document saved successfully to {output_path}!")

if __name__ == "__main__":
    main()
