# -*- coding: utf-8 -*-
import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_table_no_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def build_report_tuan_4_official():
    template_file = r"c:\thuyet minh\Docs\2024-07-Mẫu báo cáo gặp GVHD hàng tuần.docx"
    doc = Document(template_file)

    # 1. P2: BÁO CÁO TIẾN ĐỘ TUẦN 04
    p2 = doc.paragraphs[2]
    p2.text = ""
    r1 = p2.add_run("BÁO CÁO TIẾN ĐỘ TUẦN ")
    r1.font.name = "Arial"
    r1.font.size = Pt(16)
    r1.font.bold = True
    r2 = p2.add_run("04")
    r2.font.name = "Arial"
    r2.font.size = Pt(16)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(255, 0, 0)

    # P3: Thời gian
    p3 = doc.paragraphs[3]
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.text = ""
    r3 = p3.add_run("(Thời gian: 14/09/2026 – 20/09/2026)")
    r3.font.name = "Arial"
    r3.font.size = Pt(11)
    r3.font.italic = True

    # 2. Thông tin chung
    # P5: MSSV
    p5 = doc.paragraphs[5]
    p5.text = ""
    r = p5.add_run("MSSV:\t")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r = p5.add_run("123001364, 123000946")
    r.font.name = "Arial"
    r.font.size = Pt(11)

    # P6: Họ và tên
    p6 = doc.paragraphs[6]
    p6.text = ""
    r = p6.add_run("Họ và tên:\t")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r = p6.add_run("Đỗ Tấn Du, Đoàn Minh Quân")
    r.font.name = "Arial"
    r.font.size = Pt(11)

    # P7: Lớp
    p7 = doc.paragraphs[7]
    p7.text = ""
    r = p7.add_run("Lớp :\t")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r = p7.add_run("Học phần Phát triển ứng dụng - Khoa CNTT")
    r.font.name = "Arial"
    r.font.size = Pt(11)

    # P8: GVHD
    p8 = doc.paragraphs[8]
    p8.text = ""
    r = p8.add_run("GVHD:\t")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r = p8.add_run("ThS. Lê Minh Nhật")
    r.font.name = "Arial"
    r.font.size = Pt(11)

    # P9: Tên đề tài
    p9 = doc.paragraphs[9]
    p9.text = ""
    r = p9.add_run("Tên đề tài:\t")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r = p9.add_run("Nền tảng Quản trị Chuỗi cung ứng Chống lãng phí (Supply Chain Spoilage Predictor)")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True

    # 3. THÔNG TIN LÀM VIỆC
    # Danh sách công việc tuần 4
    tasks_done = [
        ("Tiếp thu & tinh gọn nghiệp vụ theo chỉ đạo của GVHD ThS. Lê Minh Nhật: ", 
         "Khoanh vùng trọng tâm vào mô hình Bán lẻ Retail (Kho tổng DC -> Cửa hàng -> Khách hàng); cơ chế xuất kho chính là bán hàng ra (Sales Out) theo nguyên tắc FEFO; cắt giảm hoàn toàn các quy trình thừa (luân chuyển ngang giữa các shop, trả hàng NCC phức tạp)."),
        ("Chuẩn hóa 3 ngưỡng an toàn khoa học theo học thuật và thực tế bán lẻ: ", 
         "1) Ngưỡng cận date theo % Vòng đời còn lại (RSL: Vàng <= 20%, Đỏ <= 10%, Tím <= 0 ngày - khóa bán); 2) Ngưỡng tồn kho tối thiểu theo mô hình ROP = d*L + SS; 3) Ngưỡng nguy cơ lãng phí Spoilage Risk (DOS > DUE: tồn kho quá nhiều so với tốc độ bán)."),
        ("Thiết kế hoàn thiện bộ sơ đồ hệ thống: ", 
         "Hoàn thành Sơ đồ Use Case (9 chức năng phân quyền Store Manager và Store Staff) và Sơ đồ luồng dữ liệu DFD Mức 0 (Context), DFD Mức 1 (5 tiến trình nghiệp vụ, 3 kho dữ liệu cốt lõi)."),
        ("Hiện thực hóa Cơ sở dữ liệu trên Microsoft SQL Server (T-SQL): ", 
         "Xây dựng 12 bảng chuẩn 3NF, 4 Views tự động cảnh báo hạn dùng và gợi ý nhập hàng động, Stored Procedures xuất kho FEFO tự động, Trigger ghi nhận lãng phí và Function fn_calculate_spoilage_risk. Nạp thành công vào SQL Server LocalDB ((localdb)\\mssqllocaldb), hiển thị font tiếng Việt mượt mà trên SSMS."),
        ("Kiểm thử tự động logic CSDL bằng Python: ", 
         "Viết kịch bản test runner (test_database_runner.py) tự động kết nối và xác minh 100% tính đúng đắn của logic tính ngày cận date và thuật toán xuất kho FEFO."),
        ("Quản lý mã nguồn & Tổ chức dự án trên Git/GitHub: ", 
         "Chuẩn hóa cấu trúc thư mục modular (Docs, Database, Sample_Data, Backend, Frontend); cấu hình .gitignore; đưa toàn bộ dự án lên GitHub (https://github.com/DoTanDu/supply-chain-spoilage-predictor) và kết nối thành viên nhóm Đoàn Minh Quân cùng tham gia.")
    ]

    # Danh sách công việc tuần 5
    tasks_next = [
        ("Khởi tạo dự án Frontend bằng React (Vite) + TailwindCSS: ", 
         "Xây dựng khung Layout Dashboard tổng thể (Sidebar, Header, Alert Badge); dựng giao diện Quản lý Lô hàng & Bảng cảnh báo hạn sử dụng (FEFO)."),
        ("Khởi tạo dự án Backend API kết nối CSDL SQL Server LocalDB: ", 
         "Viết các RESTful API xác thực người dùng (JWT) và API danh mục sản phẩm, nhập lô hàng."),
        ("Phối hợp nhóm và quản lý tiến độ: ", 
         "Phân chia module cụ thể giữa Du (Frontend) và Quân (Backend), commit và push mã nguồn thường xuyên lên GitHub.")
    ]

    # Clear placeholder paragraphs P13 to P25
    # Let's see: we can remove paragraphs from P13 down to P25, then insert properly formatted ones
    # To avoid index shifting bugs, let's collect elements or build cleanly
    # Let's inspect paragraphs list
    paras = doc.paragraphs

    # We will clear texts and repurpose or remove
    # Let's replace P13, P14, P15 with first 3 tasks, and add more paragraphs before P16
    p13 = paras[13]
    p13.text = ""
    p13.paragraph_format.line_spacing = 1.2
    p13.paragraph_format.space_after = Pt(4)
    r = p13.add_run("- " + tasks_done[0][0])
    r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True
    r = p13.add_run(tasks_done[0][1])
    r.font.name = "Arial"; r.font.size = Pt(10.5)

    p14 = paras[14]
    p14.text = ""
    p14.paragraph_format.line_spacing = 1.2
    p14.paragraph_format.space_after = Pt(4)
    r = p14.add_run("- " + tasks_done[1][0])
    r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True
    r = p14.add_run(tasks_done[1][1])
    r.font.name = "Arial"; r.font.size = Pt(10.5)

    p15 = paras[15]
    p15.text = ""
    p15.paragraph_format.line_spacing = 1.2
    p15.paragraph_format.space_after = Pt(4)
    r = p15.add_run("- " + tasks_done[2][0])
    r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True
    r = p15.add_run(tasks_done[2][1])
    r.font.name = "Arial"; r.font.size = Pt(10.5)

    # Insert tasks 3, 4, 5 before P16
    p16_elem = paras[16]._p
    for item in tasks_done[3:]:
        new_p = doc.add_paragraph()
        new_p.paragraph_format.line_spacing = 1.2
        new_p.paragraph_format.space_after = Pt(4)
        r = new_p.add_run("- " + item[0])
        r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True
        r = new_p.add_run(item[1])
        r.font.name = "Arial"; r.font.size = Pt(10.5)
        p16_elem.addprevious(new_p._p)

    # P16 header
    p16 = paras[16]
    p16.paragraph_format.space_before = Pt(8)
    p16.paragraph_format.space_after = Pt(4)

    # Tasks next week: P17, P18, P19
    p17 = paras[17]
    p17.text = ""
    p17.paragraph_format.line_spacing = 1.2
    p17.paragraph_format.space_after = Pt(4)
    r = p17.add_run("- " + tasks_next[0][0])
    r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True
    r = p17.add_run(tasks_next[0][1])
    r.font.name = "Arial"; r.font.size = Pt(10.5)

    p18 = paras[18]
    p18.text = ""
    p18.paragraph_format.line_spacing = 1.2
    p18.paragraph_format.space_after = Pt(4)
    r = p18.add_run("- " + tasks_next[1][0])
    r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True
    r = p18.add_run(tasks_next[1][1])
    r.font.name = "Arial"; r.font.size = Pt(10.5)

    p19 = paras[19]
    p19.text = ""
    p19.paragraph_format.line_spacing = 1.2
    p19.paragraph_format.space_after = Pt(4)
    r = p19.add_run("- " + tasks_next[2][0])
    r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True
    r = p19.add_run(tasks_next[2][1])
    r.font.name = "Arial"; r.font.size = Pt(10.5)

    # Signatures: We replace P23, P24, P25 with our 3-person signature block
    # Clear P20, P21, P22, P23, P24, P25
    for p_idx in [20, 21, 22, 23, 24, 25]:
        paras[p_idx].text = ""

    # Add Date line before table
    p23 = paras[23]
    p23.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p23.paragraph_format.space_before = Pt(12)
    p23.paragraph_format.space_after = Pt(4)
    r_d = p23.add_run("Đồng Nai, Ngày 19 tháng 09 năm 2026")
    r_d.font.name = "Arial"
    r_d.font.size = Pt(11)
    r_d.font.italic = True

    # Insert 3-column signature table right after P23
    sig_tbl = doc.add_table(rows=2, cols=3)
    sig_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_tbl.autofit = False
    set_table_no_borders(sig_tbl)

    col_widths = [2.25, 2.25, 2.2]
    headers_sig = [
        [("XÁC NHẬN CỦA GVHD", True), ("\n(kí tên và ghi rõ họ tên)", False, True)],
        [("SINH VIÊN THỰC HIỆN", True), ("\n(kí tên và ghi rõ họ tên)", False, True)],
        [("SINH VIÊN THỰC HIỆN", True), ("\n(kí tên và ghi rõ họ tên)", False, True)]
    ]

    for i, cell in enumerate(sig_tbl.rows[0].cells):
        cell.width = Inches(col_widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for item in headers_sig[i]:
            text, bold = item[0], item[1]
            italic = item[2] if len(item) > 2 else False
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
            r.font.bold = bold
            r.font.italic = italic

    names = [
        [("\n\n\n\n", False), ("ThS. Lê Minh Nhật", True)],
        [("\n\n\n\n", False), ("Đỗ Tấn Du", True), ("\nMSSV: 123001364", False)],
        [("\n\n\n\n", False), ("Đoàn Minh Quân", True), ("\nMSSV: 123000946", False)]
    ]

    for i, cell in enumerate(sig_tbl.rows[1].cells):
        cell.width = Inches(col_widths[i])
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for item in names[i]:
            text, bold = item[0], item[1]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(10.5)
            r.font.bold = bold

    # Move sig_tbl to before P26 ('Hình ảnh:')
    p26_elem = paras[26]._p
    p26_elem.addprevious(sig_tbl._tbl)

    # 4. Mục Hình ảnh:
    p26 = paras[26]
    p26.paragraph_format.space_before = Pt(18)
    p26.text = "Hình ảnh minh chứng kết quả tuần 4:"
    for r in p26.runs:
        r.font.name = "Arial"
        r.font.size = Pt(11.5)
        r.font.bold = True
        r.font.italic = False

    # Chèn ảnh 1: Use Case
    p_img1_title = doc.add_paragraph()
    p_img1_title.paragraph_format.space_before = Pt(8)
    p_img1_title.paragraph_format.space_after = Pt(3)
    r = p_img1_title.add_run("1. Sơ đồ Use Case tổng thể hệ thống (Phân quyền Store Manager & Store Staff)")
    r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True
    
    p_img1 = doc.add_paragraph()
    p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img1.paragraph_format.space_after = Pt(8)
    p_img1.add_run().add_picture(r"c:\thuyet minh\Docs\Diagram_Images\use_case.png", width=Inches(6.2))

    # Chèn ảnh 2: DFD 0
    p_img2_title = doc.add_paragraph()
    p_img2_title.paragraph_format.space_before = Pt(8)
    p_img2_title.paragraph_format.space_after = Pt(3)
    r = p_img2_title.add_run("2. Sơ đồ luồng dữ liệu DFD Mức 0 (Context Diagram)")
    r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True

    p_img2 = doc.add_paragraph()
    p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img2.paragraph_format.space_after = Pt(8)
    p_img2.add_run().add_picture(r"c:\thuyet minh\Docs\Diagram_Images\dfd_0.png", width=Inches(6.2))

    # Chèn ảnh 3: DFD 1
    p_img3_title = doc.add_paragraph()
    p_img3_title.paragraph_format.space_before = Pt(8)
    p_img3_title.paragraph_format.space_after = Pt(3)
    r = p_img3_title.add_run("3. Sơ đồ luồng dữ liệu DFD Mức 1 (5 Tiến trình nghiệp vụ & 3 Kho dữ liệu)")
    r.font.name = "Arial"; r.font.size = Pt(10.5); r.font.bold = True

    p_img3 = doc.add_paragraph()
    p_img3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img3.paragraph_format.space_after = Pt(8)
    p_img3.add_run().add_picture(r"c:\thuyet minh\Docs\Diagram_Images\dfd_1.png", width=Inches(6.2))

    # Save to BAO_CAO_TIEN_DO_TUAN_4.docx
    output_path = r"c:\thuyet minh\Docs\BAO_CAO_TIEN_DO_TUAN_4.docx"
    doc.save(output_path)
    print(f"Successfully generated official report: {output_path}")

if __name__ == "__main__":
    build_report_tuan_4_official()
